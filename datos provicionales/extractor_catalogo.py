"""
Extractor de Datos e Imágenes para el Catálogo de Atuel Gomas.

Ubicación: datos provicionales/extractor_catalogo.py
"""

import os
import re
import io
import openpyxl
from PIL import Image

EXCEL_NAME = "Para la revista del burneeee.xlsx"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, EXCEL_NAME)


def sanitize_folder_name(name: str) -> str:
    """Sanitiza el nombre de la hoja para crear una carpeta limpia en el sistema de archivos."""
    clean = name.strip()
    # Reemplazar tildes y caracteres conflictivos
    reemplazos = {
        'Á': 'A', 'É': 'E', 'Í': 'I', 'Ó': 'O', 'Ú': 'U',
        'á': 'a', 'é': 'e', 'í': 'i', 'ó': 'o', 'ú': 'u',
        'Ñ': 'N', 'ñ': 'n', '\ufffd': 'I'
    }
    for orig, dest in reemplazos.items():
        clean = clean.replace(orig, dest)
    clean = re.sub(r'[\\/*?:"<>|()]', '', clean)
    clean = re.sub(r'\s+', '_', clean).strip('_')
    return clean


def format_cell_value(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.2f}"
    return str(val).strip()


def export_sheet_to_txt(worksheet, output_txt_path: str) -> int:
    """Extrae todas las filas y columnas no vacías a un archivo TXT formateado."""
    rows_data = []
    max_cols = 0

    for row in worksheet.iter_rows(values_only=True):
        row_str = [format_cell_value(c) for c in row]
        if any(row_str):
            while row_str and not row_str[-1]:
                row_str.pop()
            rows_data.append(row_str)
            if len(row_str) > max_cols:
                max_cols = len(row_str)

    if not rows_data:
        with open(output_txt_path, "w", encoding="utf-8") as f:
            f.write("Pagina sin datos tabulares.\n")
        return 0

    for r in rows_data:
        r.extend([""] * (max_cols - len(r)))

    col_widths = [0] * max_cols
    for r in rows_data:
        for idx, cell in enumerate(r):
            col_widths[idx] = max(col_widths[idx], len(cell))

    col_widths = [min(w, 60) for w in col_widths]

    with open(output_txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== REPORTE DE ARTICULOS - PAGINA: {worksheet.title} ===\n")
        f.write(f"Total de registros detectados: {len(rows_data)}\n")
        f.write("=" * 80 + "\n\n")

        for row_idx, r in enumerate(rows_data):
            formatted_cells = [r[i].ljust(col_widths[i]) for i in range(max_cols)]
            line = " | ".join(formatted_cells)
            f.write(line + "\n")
            if row_idx == 0:
                f.write("-" * len(line) + "\n")

    return len(rows_data)


def extract_sheet_images(worksheet, output_dir: str) -> int:
    """Extrae y guarda todas las imágenes asociadas a la hoja."""
    count = 0
    if not hasattr(worksheet, "_images") or not worksheet._images:
        return 0

    for idx, img in enumerate(worksheet._images, start=1):
        try:
            img_data = img._data()
            image_stream = io.BytesIO(img_data)
            pil_img = Image.open(image_stream)

            fmt = pil_img.format.lower() if pil_img.format else "png"
            ext = "jpg" if fmt == "jpeg" else fmt

            filename = f"imagen_{idx}.{ext}"
            filepath = os.path.join(output_dir, filename)

            pil_img.save(filepath)
            count += 1
        except Exception as e:
            print(f"   [!] Error al extraer imagen {idx} en {worksheet.title}: {e}")

    return count


def process_catalog():
    print("=" * 75)
    print("INICIANDO EXTRACCION MODULAR DE ARTICULOS E IMAGENES (ATUEL GOMAS)")
    print(f"Archivo origen: {EXCEL_PATH}")
    print("=" * 75)

    if not os.path.exists(EXCEL_PATH):
        print(f"Error: No se encontro el archivo '{EXCEL_PATH}'")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    total_sheets = len(wb.sheetnames)
    print(f"Se detectaron {total_sheets} paginas/hojas en el Excel.\n")

    resumen = []

    for index, sheet_name in enumerate(wb.sheetnames, start=1):
        clean_name = sanitize_folder_name(sheet_name)
        folder_name = f"{index:02d}_{clean_name}"
        page_dir = os.path.join(BASE_DIR, folder_name)

        os.makedirs(page_dir, exist_ok=True)

        ws = wb[sheet_name]
        txt_path = os.path.join(page_dir, "articulos.txt")

        # 1. Exportar datos de artículos a TXT
        total_rows = export_sheet_to_txt(ws, txt_path)

        # 2. Extraer imágenes asociadas a la hoja
        total_imgs = extract_sheet_images(ws, page_dir)

        print(f"[{index:02d}/{total_sheets:02d}] {sheet_name}")
        print(f"   Carpeta: {folder_name}")
        print(f"   Registros: {total_rows} filas guardadas en 'articulos.txt'")
        print(f"   Imagenes extraidas: {total_imgs}")
        print("-" * 75)

        resumen.append({
            "hoja": sheet_name,
            "carpeta": folder_name,
            "filas": total_rows,
            "imagenes": total_imgs
        })

    print("\nEXTRACCION FINALIZADA CON EXITO.")
    total_items = sum(r["filas"] for r in resumen)
    total_fotos = sum(r["imagenes"] for r in resumen)
    print(f"Total de registros de articulos leidos: {total_items:,}")
    print(f"Total de imagenes tecnicas recuperadas: {total_fotos}")


if __name__ == "__main__":
    process_catalog()
