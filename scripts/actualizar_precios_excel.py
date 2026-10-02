#!/usr/bin/env python3
"""
Script de Actualización Masiva de Precios desde Planilla Excel
Atuel Gomas - Data Specialist (ETL & Sincronización)

Permite actualizar 'precio_mayorista_b2b' y 'precio_base' de los productos en la base
de datos (PostgreSQL Supabase o SQLite local) a partir de listas de precios en formato Excel (.xlsx).

Características:
- Detección automática y flexible de columnas (código/SKU, descripción, precio neto/sin IVA).
- Soporte para hojas simples y hojas con múltiples tablas paralelas.
- Búsqueda en memoria de alta velocidad O(1) por SKU, código OEM y coincidencia de nombre.
- Actualización en bloques (batches) controlados para evitar saturación de red/memoria.
- Reporte detallado de artículos no encontrados en 'datos provicionales/reporte_articulos_no_encontrados.txt'.
- Modo simulación (--dry-run) para verificar antes de impactar en producción.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import os
import re
import sys
import time
from typing import Any

# Asegurar path del proyecto
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from sqlalchemy import create_engine, select, text, update
from sqlalchemy.orm import Session

from src.infrastructure.database.connection import (
    DATABASE_URL as DEFAULT_DB_URL,
    SessionLocal,
    engine as default_engine,
)
from src.infrastructure.database.models import ProductoModel


# ==============================================================================
# Utilidades de Normalización y Limpieza
# ==============================================================================

def clean_str(val: Any) -> str:
    """Limpia cadenas, normaliza espacios y caracteres conflictivos."""
    if val is None:
        return ""
    s = str(val).strip()
    s = s.replace("\ufffd", "")
    return re.sub(r"\s+", " ", s).strip()


def parse_price(val: Any) -> float | None:
    """Parsea importes monetarios a float ignorando signos, espacios y separadores."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val) if val > 0 else None
    s = clean_str(val).replace("$", "").replace(" ", "").replace(",", "")
    try:
        f = float(s)
        return f if f > 0 else None
    except ValueError:
        return None


def normalize_lookup(text_val: str) -> str:
    """Normaliza texto para comparación flexible (sin acentos, mayúsculas ni signos)."""
    t = text_val.strip().lower()
    t = re.sub(r"[áäàâ]", "a", t)
    t = re.sub(r"[éëèê]", "e", t)
    t = re.sub(r"[íïìî]", "i", t)
    t = re.sub(r"[óöòô]", "o", t)
    t = re.sub(r"[úüùû]", "u", t)
    t = re.sub(r"[ñ]", "n", t)
    t = re.sub(r"[^a-z0-9]", "", t)
    return t


# ==============================================================================
# Extracción de Datos de Hojas Excel
# ==============================================================================

class ExcelPriceRow:
    def __init__(
        self,
        sheet_name: str,
        row_num: int,
        code: str,
        description: str,
        price_neto: float,
        price_base: float | None = None,
    ):
        self.sheet_name = sheet_name
        self.row_num = row_num
        self.code = code
        self.description = description
        self.price_neto = price_neto
        self.price_base = price_base

    def __repr__(self) -> str:
        return f"<ExcelPriceRow {self.sheet_name}:L{self.row_num} [{self.code}] {self.description[:30]} ${self.price_neto:,.2f}>"


def parse_worksheet_rows(ws: Any, sheet_name: str) -> list[ExcelPriceRow]:
    """
    Analiza una hoja de cálculo, detecta columnas de precios y extrae las filas.
    Soporta tablas estándar y tablas paralelas (ej. mangueras moldeadas y por metro).
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # 1. Identificar columnas que contienen encabezados o valores de precios
    price_col_indices: list[int] = []
    max_cols = max(len(r) for r in rows[:15]) if rows else 0

    for c in range(max_cols):
        is_price_col = False
        for r_idx in range(min(10, len(rows))):
            r = rows[r_idx]
            if c < len(r):
                val_txt = clean_str(r[c]).lower()
                if any(k in val_txt for k in ["precio", "dólar", "dolar", "neto", "sin iva", "pvp", "lista"]):
                    # Evitar columnas de tipo cambio o cotización general si están aisladas
                    if "dólar" in val_txt and r_idx < 4 and c == 0:
                        continue
                    is_price_col = True
                    break
        if is_price_col:
            price_col_indices.append(c)

    # Si no hubo encabezados explícitos, buscar columnas con números flotantes típicos de precios
    if not price_col_indices:
        for c in range(max_cols):
            num_count = 0
            for r_idx in range(min(25, len(rows))):
                r = rows[r_idx]
                if c < len(r) and parse_price(r[c]) is not None:
                    num_count += 1
            if num_count >= 2:
                price_col_indices.append(c)

    price_col_indices = sorted(list(set(price_col_indices)))
    extracted_rows: list[ExcelPriceRow] = []

    if not price_col_indices:
        return []

    # 2. Extraer datos para cada columna de precio detectada
    for row_num, r in enumerate(rows, start=1):
        # Omitir filas que sean evidentemente encabezados
        row_txt = " ".join(clean_str(cell).lower() for cell in r if cell is not None)
        if any(h in row_txt for h in ["descripción art", "precio sin iva", "lista de precios"]):
            continue

        for p_idx in price_col_indices:
            if p_idx < len(r):
                price = parse_price(r[p_idx])
                if price is not None and price > 0:
                    # Determinar el bloque de columnas que precede a este precio
                    prev_prices = [x for x in price_col_indices if x < p_idx]
                    start_idx = (prev_prices[-1] + 1) if prev_prices else 0

                    preceding_cells = [clean_str(r[c]) for c in range(start_idx, p_idx) if clean_str(r[c])]

                    code = ""
                    desc = ""

                    if len(preceding_cells) == 1:
                        desc = preceding_cells[0]
                    elif len(preceding_cells) == 2:
                        code = preceding_cells[0]
                        desc = preceding_cells[1]
                    elif len(preceding_cells) >= 3:
                        code = preceding_cells[0]
                        desc = " - ".join(preceding_cells[1:])

                    # Extraer código de la descripción si viene en formato "ARTICULO - CODIGO"
                    if (not code or code.upper() in ("S/C", "S / C")) and " - " in desc:
                        parts = desc.split(" - ", 1)
                        if len(parts[1].strip()) <= 50 and not any(k in parts[1].lower() for k in ["rollo", "tramo", "bolsa"]):
                            code = parts[1].strip()

                    if code or desc:
                        extracted_rows.append(
                            ExcelPriceRow(
                                sheet_name=sheet_name,
                                row_num=row_num,
                                code=code,
                                description=desc,
                                price_neto=price,
                            )
                        )

    return extracted_rows


# ==============================================================================
# Motor de Búsqueda y Actualización de Productos
# ==============================================================================

class PriceUpdateService:
    def __init__(
        self,
        db_engine: Any,
        markup_factor: float = 1.30,
        batch_size: int = 500,
        dry_run: bool = False,
    ):
        self.engine = db_engine
        self.markup_factor = markup_factor
        self.batch_size = batch_size
        self.dry_run = dry_run

        # Índices en memoria para búsqueda O(1)
        self.sku_index: dict[str, list[int]] = {}
        self.oem_index: dict[str, list[int]] = {}
        self.name_exact_index: dict[str, list[int]] = {}
        self.name_norm_index: dict[str, list[int]] = {}

        self._build_product_index()

    def _build_product_index(self):
        """Carga y construye los índices de búsqueda en memoria."""
        print("\n[+] Construyendo índices de productos en memoria desde la base de datos...")
        t0 = time.time()
        with self.engine.connect() as conn:
            query = select(
                ProductoModel.id,
                ProductoModel.sku,
                ProductoModel.codigo_oem,
                ProductoModel.nombre,
                ProductoModel.precio_mayorista_b2b,
                ProductoModel.precio_base,
            )
            rows = conn.execute(query).fetchall()

        for pid, sku, oem, nombre, b2b, base in rows:
            if sku:
                s_key = clean_str(sku).upper()
                self.sku_index.setdefault(s_key, []).append(pid)
            if oem:
                o_key = clean_str(oem).upper()
                self.oem_index.setdefault(o_key, []).append(pid)
            if nombre:
                n_raw = clean_str(nombre).lower()
                self.name_exact_index.setdefault(n_raw, []).append(pid)
                n_norm = normalize_lookup(nombre)
                if n_norm:
                    self.name_norm_index.setdefault(n_norm, []).append(pid)

        duration = time.time() - t0
        print(f"    [OK] Indexados {len(rows):,} productos en {duration:.2f}s:")
        print(f"         • Claves SKU únicas:  {len(self.sku_index):,}")
        print(f"         • Claves OEM únicas:  {len(self.oem_index):,}")
        print(f"         • Nombres indexados:  {len(self.name_exact_index):,}")

    def find_product_ids(self, code: str, desc: str) -> list[int]:
        """Busca coincidencias en la base de datos por SKU, OEM o nombre."""
        # 1. Búsqueda por Código / SKU
        if code and code.upper() not in ("S/C", "S / C", "NONE", ""):
            c_key = clean_str(code).upper()
            if c_key in self.sku_index:
                return self.sku_index[c_key]
            if c_key in self.oem_index:
                return self.oem_index[c_key]

        # 2. Búsqueda por coincidencia exacta de nombre
        if desc:
            d_raw = clean_str(desc).lower()
            if d_raw in self.name_exact_index:
                return self.name_exact_index[d_raw]

            # Búsqueda con prefijo 'Manguera ' si aplica
            manguera_variant = f"manguera {d_raw}"
            if manguera_variant in self.name_exact_index:
                return self.name_exact_index[manguera_variant]

            # Búsqueda si el nombre tiene código embebido
            if " - " in desc:
                cand = desc.split(" - ", 1)[1].strip().upper()
                if cand in self.sku_index:
                    return self.sku_index[cand]
                if cand in self.oem_index:
                    return self.oem_index[cand]

            # 3. Búsqueda por texto normalizado
            d_norm = normalize_lookup(desc)
            if d_norm and d_norm in self.name_norm_index:
                return self.name_norm_index[d_norm]

        return []

    def process_excel(
        self,
        excel_path: str,
        sheet_name_filter: str | None = None,
        report_path: str | None = None,
    ) -> dict[str, Any]:
        """Procesa el archivo Excel, ejecuta actualizaciones y genera reporte de pendientes."""
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"No se encontró el archivo Excel: {excel_path}")

        print("\n" + "=" * 80)
        print(f"PROCESANDO ARCHIVO EXCEL: {excel_path}")
        if self.dry_run:
            print(">>> MODO SIMULACIÓN (DRY-RUN) ACTIVO: No se modificarán datos en la BD <<<")
        print("=" * 80)

        t_start = time.time()
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
        sheets = wb.sheetnames

        if sheet_name_filter:
            if sheet_name_filter in sheets:
                sheets = [sheet_name_filter]
            else:
                try:
                    s_idx = int(sheet_name_filter) - 1
                    sheets = [sheets[s_idx]]
                except (ValueError, IndexError):
                    raise ValueError(f"Hoja '{sheet_name_filter}' no encontrada en {excel_path}")

        total_rows_read = 0
        matched_rows_count = 0
        unmatched_items: list[dict[str, Any]] = []
        updated_product_ids: set[int] = set()

        # Diccionario para acumular los últimos precios por cada product_id
        # {product_id: {'id': pid, 'precio_mayorista_b2b': x, 'precio_base': y, 'updated_at': now}}
        pending_db_updates: dict[int, dict[str, Any]] = {}

        now_utc = datetime.now(timezone.utc)

        for sname in sheets:
            # Omitir hojas de índice general si existen
            if "INDICE" in sname.upper() and len(sheets) > 1:
                continue

            ws = wb[sname]
            extracted = parse_worksheet_rows(ws, sname)
            total_rows_read += len(extracted)

            for item in extracted:
                pids = self.find_product_ids(item.code, item.description)

                if pids:
                    matched_rows_count += 1
                    updated_product_ids.update(pids)

                    # Calcular precio base si no vino en la planilla
                    new_b2b = round(item.price_neto, 2)
                    new_base = (
                        round(item.price_base, 2)
                        if item.price_base is not None
                        else round(new_b2b * self.markup_factor, 2)
                    )

                    for pid in pids:
                        pending_db_updates[pid] = {
                            "id": pid,
                            "precio_mayorista_b2b": new_b2b,
                            "precio_base": new_base,
                            "updated_at": now_utc,
                        }
                else:
                    unmatched_items.append({
                        "sheet": item.sheet_name,
                        "row": item.row_num,
                        "code": item.code or "S/C",
                        "description": item.description or "(Sin descripción)",
                        "price": item.price_neto,
                        "reason": "No coincide con ningún SKU, OEM ni Nombre en el catálogo",
                    })

        wb.close()

        # 3. Aplicar actualizaciones en la base de datos (por lotes)
        total_products_updated = len(pending_db_updates)

        if not self.dry_run and pending_db_updates:
            print(f"\n[+] Impactando actualizaciones en la base de datos ({total_products_updated:,} productos)...")
            t_upd = time.time()
            update_list = list(pending_db_updates.values())

            with Session(self.engine) as session:
                for i in range(0, len(update_list), self.batch_size):
                    chunk = update_list[i : i + self.batch_size]
                    session.execute(update(ProductoModel), chunk)
                    if (i + len(chunk)) % 2500 == 0 or (i + len(chunk)) == len(update_list):
                        print(f"    -> Actualizados {i + len(chunk)}/{len(update_list)} productos...")
                session.commit()

            print(f"    [OK] Actualizaciones confirmadas en {time.time() - t_upd:.2f}s.")
        elif self.dry_run:
            print(f"\n[DRY-RUN] Se habrían actualizado {total_products_updated:,} productos en la BD.")

        # 4. Exportar reporte de artículos no encontrados
        if report_path and unmatched_items:
            os.makedirs(os.path.dirname(os.path.abspath(report_path)), exist_ok=True)
            with open(report_path, "w", encoding="utf-8") as f:
                f.write("=" * 110 + "\n")
                f.write("REPORTE DE ARTÍCULOS NO ENCONTRADOS EN BASE DE DATOS (ATUEL GOMAS)\n")
                f.write(f"Fecha de ejecución: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}\n")
                f.write(f"Archivo procesado:  {excel_path}\n")
                f.write(f"Total registros no encontrados: {len(unmatched_items):,}\n")
                f.write("=" * 110 + "\n\n")

                headers = ["HOJA", "FILA", "CÓDIGO / SKU", "DESCRIPCIÓN", "PRECIO LEÍDO", "MOTIVO"]
                col_w = [25, 7, 18, 45, 15, 30]
                hdr_line = " | ".join(h.ljust(w) for h, w in zip(headers, col_w))
                f.write(hdr_line + "\n")
                f.write("-" * len(hdr_line) + "\n")

                for u in unmatched_items:
                    row_data = [
                        u["sheet"][:col_w[0]],
                        str(u["row"]),
                        u["code"][:col_w[2]],
                        u["description"][:col_w[3]],
                        f"${u['price']:,.2f}",
                        u["reason"][:col_w[5]],
                    ]
                    line = " | ".join(val.ljust(w) for val, w in zip(row_data, col_w))
                    f.write(line + "\n")

            print(f"\n[!] Reporte de artículos no encontrados exportado a:\n    {report_path}")

        total_duration = time.time() - t_start

        # 5. Imprimir balance claro
        print("\n" + "=" * 80)
        print("BALANCE DE ACTUALIZACIÓN DE PRECIOS")
        print("=" * 80)
        print(f"  • Total filas leídas del Excel:        {total_rows_read:>6,}")
        print(f"  • Filas con coincidencia (match):      {matched_rows_count:>6,}")
        print(f"  • Productos actualizados con éxito:    {total_products_updated:>6,}")
        print(f"  • Filas no encontradas (reportadas):   {len(unmatched_items):>6,}")
        print(f"  • Tiempo total de ejecución:           {total_duration:>6.2f} segundos")
        print("=" * 80)

        return {
            "total_rows_read": total_rows_read,
            "matched_rows_count": matched_rows_count,
            "total_products_updated": total_products_updated,
            "unmatched_count": len(unmatched_items),
            "report_path": report_path if unmatched_items else None,
            "duration": total_duration,
        }


# ==============================================================================
# Punto de Entrada
# ==============================================================================

def main():
    parser = argparse.ArgumentParser(description="Actualización masiva de precios desde Excel para Atuel Gomas")
    parser.add_argument(
        "excel_path",
        nargs="?",
        default=os.path.join(PROJECT_ROOT, "datos provicionales", "Para la revista del burneeee.xlsx"),
        help="Ruta al archivo Excel (.xlsx) con la lista de precios",
    )
    parser.add_argument("--sheet", default=None, help="Nombre o índice de la hoja a procesar (por defecto: todas)")
    parser.add_argument("--markup", type=float, default=1.30, help="Factor de recargo para precio base (def: 1.30 = +30%%)")
    parser.add_argument("--batch-size", type=int, default=500, help="Lote de actualización SQL (def: 500)")
    parser.add_argument("--dry-run", action="store_true", help="Modo simulación: no guarda cambios en la base de datos")
    parser.add_argument("--sqlite", action="store_true", help="Forzar conexión a base local SQLite 'atuel_gomas.db'")
    parser.add_argument("--db-url", default=None, help="Cadena de conexión personalizada")
    parser.add_argument(
        "--report-file",
        default=os.path.join(PROJECT_ROOT, "datos provicionales", "reporte_articulos_no_encontrados.txt"),
        help="Ruta para el archivo de reporte de artículos no encontrados",
    )

    args = parser.parse_args()

    # Determinar engine de base de datos
    if args.sqlite:
        target_url = f"sqlite:///{os.path.join(PROJECT_ROOT, 'atuel_gomas.db')}"
        active_engine = create_engine(target_url)
        db_label = f"SQLite local ({target_url})"
    elif args.db_url:
        active_engine = create_engine(args.db_url)
        db_label = f"Personalizada ({args.db_url[:25]}...)"
    else:
        active_engine = default_engine
        db_label = f"Producción / Configurada ({DEFAULT_DB_URL[:30]}...)"

    print("=" * 80)
    print("ATUEL GOMAS - ACTUALIZADOR DE PRECIOS MASIVO DESDE EXCEL")
    print("=" * 80)
    print(f"Base de datos destino: {db_label}")
    print(f"Factor de recargo PVP: x{args.markup:.2f} (+{(args.markup - 1)*100:.0f}%)")

    service = PriceUpdateService(
        db_engine=active_engine,
        markup_factor=args.markup,
        batch_size=args.batch_size,
        dry_run=args.dry_run,
    )

    service.process_excel(
        excel_path=args.excel_path,
        sheet_name_filter=args.sheet,
        report_path=args.report_file,
    )


if __name__ == "__main__":
    main()
